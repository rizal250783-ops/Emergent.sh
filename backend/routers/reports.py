import io
import os

from bson import ObjectId
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle, Paragraph,
                                Spacer, Image)
from reportlab.graphics.shapes import Drawing, String
from reportlab.graphics.charts.linecharts import HorizontalLineChart
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.charts.textlabels import Label

from database import db
from security import get_current_user, require_roles
from calc import build_kpis, build_riwayat, build_leaderboard, compute_achievement

router = APIRouter(prefix="/reports", tags=["reports"])

EMERALD = colors.HexColor("#047857")
EMERALD_DARK = colors.HexColor("#064E3B")
GOLD = colors.HexColor("#D97706")
LIGHT = colors.HexColor("#ECFDF5")
LOGO_PATH = "/app/frontend/public/logo.png"

STATUS_ID = {"excellent": "Sangat Baik", "good": "Baik", "need_attention": "Perlu Perhatian", "critical": "Kritis", "na": "N/A"}
MONTHS = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober", "November", "Desember"]


def rupiah(n):
    if n is None:
        return "Rp 0"
    return "Rp " + f"{int(n):,}".replace(",", ".")


def periode_label(p):
    if not p:
        return ""
    y, m = p.split("-")
    return f"{MONTHS[int(m) - 1]} {y}"


def _short_month(p):
    return MONTHS[int(p.split("-")[1]) - 1][:3]


def build_trend_charts(riwayat):
    """Return a Drawing with achievement line chart + realisasi bar chart."""
    labels = [_short_month(r["bulan"]) for r in riwayat]
    ach = [(r["achievement"] or 0) for r in riwayat]
    real = [(r["realisasi"] or 0) for r in riwayat]

    d = Drawing(500, 190)
    # Line chart: achievement %
    lc = HorizontalLineChart()
    lc.x = 30
    lc.y = 20
    lc.width = 200
    lc.height = 140
    lc.data = [ach]
    lc.categoryAxis.categoryNames = labels
    lc.categoryAxis.labels.fontSize = 7
    lc.valueAxis.valueMin = 0
    lc.valueAxis.valueMax = max(120, (max(ach) if ach else 0) + 20)
    lc.valueAxis.labels.fontSize = 7
    lc.lines[0].strokeColor = GOLD
    lc.lines[0].strokeWidth = 2
    lc.lines.symbol = None
    d.add(String(30, 172, "Tren Achievement (%)", fontSize=9, fillColor=EMERALD_DARK, fontName="Helvetica-Bold"))
    d.add(lc)

    # Bar chart: realisasi (in juta)
    real_juta = [round(v / 1_000_000, 1) for v in real]
    bc = VerticalBarChart()
    bc.x = 300
    bc.y = 20
    bc.width = 180
    bc.height = 140
    bc.data = [real_juta]
    bc.categoryAxis.categoryNames = labels
    bc.categoryAxis.labels.fontSize = 7
    bc.valueAxis.valueMin = 0
    bc.valueAxis.labels.fontSize = 7
    bc.bars[0].fillColor = EMERALD
    bc.barWidth = 6
    d.add(String(300, 172, "Realisasi (Juta Rp)", fontSize=9, fillColor=EMERALD_DARK, fontName="Helvetica-Bold"))
    d.add(bc)
    return d


@router.get("/ao-pdf/{ao_id}")
async def ao_pdf(ao_id: str, periode: str, user=Depends(get_current_user)):
    if user["jabatan"] not in ("Admin", "Direktur") and str(user["_id"]) != ao_id:
        raise HTTPException(status_code=403, detail="Akses ditolak")
    try:
        oid = ObjectId(ao_id)
    except Exception:
        raise HTTPException(status_code=400, detail="ID AO tidak valid")
    target = await db.users.find_one({"_id": oid})
    if not target:
        raise HTTPException(status_code=404, detail="AO tidak ditemukan")

    kpis = await build_kpis(target, periode)
    riwayat = await build_riwayat(ao_id, target["jabatan"])

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=18 * mm, bottomMargin=16 * mm,
                            leftMargin=16 * mm, rightMargin=16 * mm, title=f"Rekap {target['nama']}")
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Heading1"], textColor=EMERALD_DARK, fontSize=18, spaceAfter=2)
    small = ParagraphStyle("small", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#64748b"))
    label = ParagraphStyle("label", parent=styles["Normal"], fontSize=10, textColor=colors.HexColor("#334155"))
    sec = ParagraphStyle("sec", parent=styles["Heading2"], textColor=EMERALD, fontSize=13, spaceBefore=10, spaceAfter=6)
    white = ParagraphStyle("white", parent=styles["Normal"], fontSize=9, textColor=colors.white)

    elems = []
    # Header band
    header_cells = [[]]
    logo = None
    if os.path.exists(LOGO_PATH):
        try:
            logo = Image(LOGO_PATH, width=16 * mm, height=16 * mm)
        except Exception:
            logo = ""
    title_block = [
        Paragraph("<b>PT BPRS HAJI MISKIN</b>", ParagraphStyle("t", parent=styles["Normal"], fontSize=13, textColor=colors.white)),
        Paragraph("AO-360 · Laporan Rekap Pencapaian AO", white),
    ]
    header = Table([[logo or "", title_block]], colWidths=[20 * mm, 150 * mm])
    header.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), EMERALD_DARK),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    elems.append(header)
    elems.append(Spacer(1, 10))

    # AO info
    info = Table([
        [Paragraph("<b>Nama AO</b>", label), Paragraph(target["nama"], label),
         Paragraph("<b>Periode</b>", label), Paragraph(periode_label(periode), label)],
        [Paragraph("<b>Kode Marketing</b>", label), Paragraph(target["kode_marketing"], label),
         Paragraph("<b>Jabatan</b>", label), Paragraph(target["jabatan"], label)],
    ], colWidths=[35 * mm, 50 * mm, 30 * mm, 55 * mm])
    info.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), LIGHT),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#A7F3D0")),
        ("INNERGRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D1FAE5")),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
    ]))
    elems.append(info)

    # KPI table
    elems.append(Paragraph("Ringkasan Pencapaian", sec))
    kpi_head = ["Komponen", "Target", "Realisasi", "Achievement", "Status"]
    kpi_rows = [kpi_head]
    for k in kpis:
        ach = "N/A" if k["achievement"] is None else f"{k['achievement']:.1f}%"
        kpi_rows.append([k["komponen"], rupiah(k["target"]), rupiah(k["realisasi"]), ach, STATUS_ID.get(k["status"], "-")])
    if len(kpi_rows) == 1:
        kpi_rows.append(["Tidak ada komponen", "-", "-", "-", "-"])
    kt = Table(kpi_rows, colWidths=[40 * mm, 35 * mm, 35 * mm, 28 * mm, 32 * mm])
    kt.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), EMERALD),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#E2E8F0")),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("ALIGN", (1, 0), (-1, -1), "CENTER"),
    ]))
    elems.append(kt)

    # Trend charts
    if riwayat:
        elems.append(Paragraph("Grafik Tren Performa", sec))
        try:
            elems.append(build_trend_charts(riwayat))
        except Exception:
            pass

    # Riwayat table
    elems.append(Paragraph("Riwayat Performance Bulanan", sec))
    rhead = ["Bulan", "Target", "Realisasi", "Achievement", "Status"]
    rrows = [rhead]
    for r in riwayat:
        ach = "N/A" if r["achievement"] is None else f"{r['achievement']:.1f}%"
        rrows.append([periode_label(r["bulan"]), rupiah(r["target"]), rupiah(r["realisasi"]), ach, STATUS_ID.get(r["status"], "-")])
    rt = Table(rrows, colWidths=[38 * mm, 37 * mm, 37 * mm, 28 * mm, 30 * mm])
    rt.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), EMERALD_DARK),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#E2E8F0")),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("ALIGN", (1, 0), (-1, -1), "CENTER"),
    ]))
    elems.append(rt)

    elems.append(Spacer(1, 14))
    elems.append(Paragraph(f"Dicetak oleh {user['nama']} ({user['jabatan']}) · AO-360 PT BPRS Haji Miskin", small))

    doc.build(elems)
    buf.seek(0)
    fname = f"Rekap_{target['kode_marketing']}_{periode}.pdf"
    return StreamingResponse(buf, media_type="application/pdf",
                             headers={"Content-Disposition": f"attachment; filename={fname}"})



def _summary_table(title, komponen, rows):
    head = ["#", "Nama", "Kode", "Target", "Realisasi", "Ach.", "Status"]
    data = [head]
    for r in rows:
        ach = "N/A" if r["achievement"] is None else f"{r['achievement']:.0f}%"
        data.append([str(r.get("ranking", "-")), r["nama"], r["kode_marketing"],
                     rupiah(r["target"]), rupiah(r["realisasi"]), ach, STATUS_ID.get(r["status"], "-")])
    if len(data) == 1:
        data.append(["-", "Belum ada data", "-", "-", "-", "-", "-"])
    t = Table(data, colWidths=[8 * mm, 42 * mm, 16 * mm, 34 * mm, 34 * mm, 16 * mm, 28 * mm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), EMERALD),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#E2E8F0")),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("ALIGN", (3, 0), (-1, -1), "CENTER"),
    ]))
    return t


@router.get("/team-pdf")
async def team_pdf(periode: str, user=Depends(require_roles("Direktur", "Admin"))):
    styles = getSampleStyleSheet()
    sec = ParagraphStyle("sec", parent=styles["Heading2"], textColor=EMERALD, fontSize=13, spaceBefore=10, spaceAfter=6)
    small = ParagraphStyle("small", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#64748b"))
    white = ParagraphStyle("white", parent=styles["Normal"], fontSize=9, textColor=colors.white)

    total_pencairan = sum(d["jumlah_pencairan"] for d in await db.lending_achievement_details.find({"periode": periode}).to_list(5000))
    total_funding = sum(d["jumlah_simpanan"] for d in await db.funding_achievement_details.find({"periode": periode}).to_list(5000))
    rec_docs = await db.recovery_achievement_details.find({"periode": periode}).to_list(5000)
    total_recovery = sum(d["jumlah_recovery"] for d in rec_docs if d.get("kolektibilitas") == 3 and not d.get("is_write_off"))
    targets = await db.targets.find({"periode": periode}).to_list(500)
    tp = sum(t.get("target_pencairan", 0) or 0 for t in targets)
    tf = sum(t.get("target_funding", 0) or 0 for t in targets)
    tr = sum(t.get("target_recovery", 0) or 0 for t in targets)

    lb_p = await build_leaderboard("Pembiayaan", periode)
    lb_f = await build_leaderboard("Funding", periode)
    lb_r = await build_leaderboard("Recovery", periode)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=16 * mm, bottomMargin=14 * mm,
                            leftMargin=14 * mm, rightMargin=14 * mm, title=f"Laporan Tim {periode}")
    elems = []
    logo = Image(LOGO_PATH, width=16 * mm, height=16 * mm) if os.path.exists(LOGO_PATH) else ""
    header = Table([[logo, [
        Paragraph("<b>PT BPRS HAJI MISKIN</b>", ParagraphStyle("t", parent=styles["Normal"], fontSize=13, textColor=colors.white)),
        Paragraph(f"AO-360 · Laporan Tim (Rapat Manajemen) · {periode_label(periode)}", white),
    ]]], colWidths=[20 * mm, 162 * mm])
    header.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), EMERALD_DARK), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                                ("LEFTPADDING", (0, 0), (-1, -1), 8), ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
    elems.append(header)
    elems.append(Spacer(1, 10))

    # bank-wide summary
    elems.append(Paragraph("Ringkasan Bank-Wide", sec))
    def line(nm, real, tgt):
        r = compute_achievement(real, tgt)
        ach = "N/A" if r["achievement"] is None else f"{r['achievement']:.0f}%"
        return [nm, rupiah(tgt), rupiah(real), ach, STATUS_ID.get(r["status"], "-")]
    st = Table([["Komponen", "Total Target", "Total Realisasi", "Achievement", "Status"],
                line("Pembiayaan", total_pencairan, tp), line("Funding", total_funding, tf), line("Recovery Kol.3", total_recovery, tr)],
               colWidths=[45 * mm, 40 * mm, 40 * mm, 28 * mm, 29 * mm])
    st.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), EMERALD), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 9),
                            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#E2E8F0")),
                            ("ALIGN", (1, 0), (-1, -1), "CENTER"), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    elems.append(st)

    elems.append(Paragraph("Ranking Pembiayaan", sec)); elems.append(_summary_table("Pembiayaan", "Pembiayaan", lb_p))
    elems.append(Paragraph("Ranking Funding", sec)); elems.append(_summary_table("Funding", "Funding", lb_f))
    elems.append(Paragraph("Ranking Recovery (Kol.3)", sec)); elems.append(_summary_table("Recovery", "Recovery", lb_r))
    elems.append(Spacer(1, 12))
    elems.append(Paragraph(f"Dicetak oleh {user['nama']} ({user['jabatan']}) · AO-360 PT BPRS Haji Miskin", small))

    doc.build(elems)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/pdf",
                             headers={"Content-Disposition": f"attachment; filename=Laporan_Tim_{periode}.pdf"})


@router.get("/ao-excel/{ao_id}")
async def ao_excel(ao_id: str, periode: str, user=Depends(get_current_user)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    if user["jabatan"] not in ("Admin", "Direktur") and str(user["_id"]) != ao_id:
        raise HTTPException(status_code=403, detail="Akses ditolak")
    try:
        oid = ObjectId(ao_id)
    except Exception:
        raise HTTPException(status_code=400, detail="ID AO tidak valid")
    target = await db.users.find_one({"_id": oid})
    if not target:
        raise HTTPException(status_code=404, detail="AO tidak ditemukan")

    kpis = await build_kpis(target, periode)
    riwayat = await build_riwayat(ao_id, target["jabatan"])

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="047857")
    head_font = Font(bold=True, color="FFFFFF")

    ws = wb.active
    ws.title = "Ringkasan"
    ws["A1"] = "PT BPRS HAJI MISKIN — Rekap Pencapaian AO"
    ws["A1"].font = Font(bold=True, size=13, color="064E3B")
    ws["A2"] = f"Nama: {target['nama']}"
    ws["A3"] = f"Kode: {target['kode_marketing']}  |  Jabatan: {target['jabatan']}  |  Periode: {periode_label(periode)}"
    hdr = ["Komponen", "Target", "Realisasi", "Achievement (%)", "Status"]
    ws.append([])
    ws.append(hdr)
    for i, c in enumerate(hdr, 1):
        cell = ws.cell(row=5, column=i); cell.fill = head_fill; cell.font = head_font
    for k in kpis:
        ws.append([k["komponen"], k["target"], k["realisasi"],
                   ("N/A" if k["achievement"] is None else k["achievement"]), STATUS_ID.get(k["status"], "-")])
    for col in "ABCDE":
        ws.column_dimensions[col].width = 20

    ws2 = wb.create_sheet("Riwayat Bulanan")
    rhdr = ["Bulan", "Target", "Realisasi", "Achievement (%)", "Status"]
    ws2.append(rhdr)
    for i, c in enumerate(rhdr, 1):
        cell = ws2.cell(row=1, column=i); cell.fill = head_fill; cell.font = head_font
    for r in riwayat:
        ws2.append([periode_label(r["bulan"]), r["target"], r["realisasi"],
                    ("N/A" if r["achievement"] is None else r["achievement"]), STATUS_ID.get(r["status"], "-")])
    for col in "ABCDE":
        ws2.column_dimensions[col].width = 20

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"Rekap_{target['kode_marketing']}_{periode}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename={fname}"})
