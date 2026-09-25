"""BSI ASSET DEAL — Iteration 3 tests.
Covers: 25 demo assets w/ 2 SOLD sorted last, location cascade DKI totals,
batch/similar endpoints, SOLD/UNSOLD RCG flow + RBAC + approval history,
Excel report link + xlsx validation, dashboard by_status SOLD key.
"""
import io
import os
import pytest
import requests

BASE_URL = os.environ.get('EXPO_PUBLIC_BACKEND_URL', 'https://mobile-app-build-11274.preview.emergentagent.com').rstrip('/')
API = f"{BASE_URL}/api"


def _login(username, password):
    r = requests.post(f"{API}/auth/login", json={"username": username, "password": password}, timeout=30)
    assert r.status_code == 200, f"login {username}: {r.status_code} {r.text}"
    return r.json()


def H(t):
    return {"Authorization": f"Bearer {t}"}


@pytest.fixture(scope="module")
def admin_token():
    return _login("admin", "Admin@2026")["access_token"]


@pytest.fixture(scope="module")
def ma_token():
    return _login("2190020284", "BSI@2026")["access_token"]


# ---------- Catalog: 25 assets, 2 SOLD sorted last ----------
class TestCatalogSoldSort:
    def test_total_and_sold_at_end(self):
        r = requests.get(f"{API}/public/catalog", params={"limit": 30}, timeout=20)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["total"] == 25, f"expected total=25, got {d['total']}"
        items = d["items"]
        assert len(items) == 25
        # Last two must be SOLD, none before them
        sold_flags = [it.get("is_sold") for it in items]
        assert sold_flags[-2:] == [True, True], f"last 2 must be SOLD, got {sold_flags[-2:]}"
        assert not any(sold_flags[:-2]), "SOLD items must be at the end only"
        # SOLD items hide pic_wa; non-sold expose pic_wa
        for it in items[:-2]:
            assert it.get("pic_wa"), f"non-sold missing pic_wa: {it.get('id')}"
            assert it.get("is_sold") is False
        for it in items[-2:]:
            assert it.get("pic_wa") in (None, ""), f"SOLD leaked pic_wa: {it}"
            assert it.get("status") == "SOLD"

    def test_dki_jakarta_cascade_counts(self):
        r = requests.get(f"{API}/public/catalog", params={"provinsi": "DKI Jakarta"}, timeout=15)
        assert r.status_code == 200
        assert r.json()["total"] == 5, f"DKI expected 5, got {r.json()['total']}"

        r = requests.get(f"{API}/public/catalog",
                         params={"provinsi": "DKI Jakarta", "kabupaten_kota": "Kota Jakarta Selatan"},
                         timeout=15)
        assert r.status_code == 200
        assert r.json()["total"] == 2, f"DKI/JakSel expected 2, got {r.json()['total']}"

        r = requests.get(f"{API}/public/catalog",
                         params={"provinsi": "DKI Jakarta",
                                 "kabupaten_kota": "Kota Jakarta Selatan",
                                 "kecamatan": "Cilandak"},
                         timeout=15)
        assert r.status_code == 200
        assert r.json()["total"] == 1, f"DKI/JakSel/Cilandak expected 1, got {r.json()['total']}"


# ---------- Batch and Similar ----------
class TestBatchAndSimilar:
    def test_batch_valid_and_bogus_ordering(self):
        cat = requests.get(f"{API}/public/catalog", params={"limit": 30}, timeout=15).json()
        # pick two non-sold ids (from first two items)
        id1 = cat["items"][0]["id"]
        id2 = cat["items"][1]["id"]
        r = requests.get(f"{API}/public/catalog/batch",
                         params={"ids": f"{id1},{id2},bogus"}, timeout=15)
        assert r.status_code == 200, r.text
        out = r.json()
        assert isinstance(out, list) and len(out) == 2
        assert out[0]["id"] == id1
        assert out[1]["id"] == id2

    def test_similar_endpoint(self):
        cat = requests.get(f"{API}/public/catalog", params={"limit": 30}, timeout=15).json()
        # first (non-sold) item
        aid = cat["items"][0]["id"]
        r = requests.get(f"{API}/public/catalog/{aid}/similar", timeout=15)
        assert r.status_code == 200
        arr = r.json()
        assert isinstance(arr, list)
        assert len(arr) <= 6
        for it in arr:
            assert it["id"] != aid
            assert it.get("status") != "SOLD"
            assert not it.get("is_sold"), f"sold should be excluded: {it}"


# ---------- SOLD / UNSOLD flow ----------
class TestSoldFlow:
    """Uses one PUBLISHED asset, marks SOLD then UNSOLD back (cleanup)."""

    @pytest.fixture(scope="class", autouse=True)
    def target_asset(self, admin_token):
        r = requests.get(f"{API}/public/catalog", params={"limit": 30}, timeout=15).json()
        # first non-sold from list
        non_sold = [it for it in r["items"] if not it.get("is_sold")]
        assert non_sold, "no non-sold in catalog"
        pytest.sold_aid = non_sold[0]["id"]
        yield
        # cleanup: if left SOLD, unsold it
        det = requests.get(f"{API}/public/catalog/{pytest.sold_aid}", timeout=10)
        if det.status_code == 200 and det.json().get("is_sold"):
            requests.post(f"{API}/rcg/assets/{pytest.sold_aid}/unsold",
                          headers=H(admin_token), timeout=15)

    def test_ma_cannot_sold(self, ma_token):
        r = requests.post(f"{API}/rcg/assets/{pytest.sold_aid}/sold",
                          headers=H(ma_token), json={"notes": "tes"}, timeout=15)
        assert r.status_code == 403

    def test_admin_mark_sold(self, admin_token):
        r = requests.post(f"{API}/rcg/assets/{pytest.sold_aid}/sold",
                          headers=H(admin_token), json={"notes": "tes"}, timeout=15)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d.get("ok") is True
        assert d.get("status") == "SOLD"

    def test_public_detail_after_sold(self):
        r = requests.get(f"{API}/public/catalog/{pytest.sold_aid}", timeout=15)
        assert r.status_code == 200
        d = r.json()
        assert d.get("is_sold") is True
        assert d.get("status") == "SOLD"
        assert d.get("pic_wa") in (None, ""), f"pic_wa leaked when SOLD: {d.get('pic_wa')}"

    def test_double_sold_409(self, admin_token):
        r = requests.post(f"{API}/rcg/assets/{pytest.sold_aid}/sold",
                          headers=H(admin_token), json={"notes": "again"}, timeout=15)
        assert r.status_code == 409, r.text

    def test_unsold(self, admin_token):
        r = requests.post(f"{API}/rcg/assets/{pytest.sold_aid}/unsold",
                          headers=H(admin_token), timeout=15)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d.get("status") == "PUBLISHED"

    def test_public_after_unsold(self):
        r = requests.get(f"{API}/public/catalog/{pytest.sold_aid}", timeout=15)
        assert r.status_code == 200
        d = r.json()
        assert d.get("status") == "PUBLISHED"
        assert d.get("is_sold") in (False, None)
        assert d.get("pic_wa"), "pic_wa should be present again after unsold"

    def test_approval_history_has_sold_actions(self, admin_token):
        r = requests.get(f"{API}/assets/{pytest.sold_aid}",
                         headers=H(admin_token), timeout=15)
        assert r.status_code == 200
        det = r.json()
        hist = det.get("approval_history") or []
        actions = [h.get("action") for h in hist]
        assert "MARK_SOLD" in actions, f"MARK_SOLD missing in history: {actions}"
        assert "UNMARK_SOLD" in actions, f"UNMARK_SOLD missing in history: {actions}"


# ---------- Excel report ----------
class TestExcelReport:
    def test_ma_forbidden(self, ma_token):
        r = requests.get(f"{API}/rcg/reports/assets/link", headers=H(ma_token), timeout=15)
        assert r.status_code == 403

    def test_link_and_download(self, admin_token):
        r = requests.get(f"{API}/rcg/reports/assets/link",
                         headers=H(admin_token), timeout=15)
        assert r.status_code == 200, r.text
        j = r.json()
        assert j["url"].startswith("/api/rcg/reports/assets.xlsx?token="), j["url"]

        r2 = requests.get(f"{BASE_URL}{j['url']}", timeout=30)
        assert r2.status_code == 200, r2.text[:200]
        ct = r2.headers.get("content-type", "")
        assert "spreadsheet" in ct or "officedocument" in ct, f"bad content-type: {ct}"
        # xlsx magic (zip)
        assert r2.content[:2] == b"PK", "not a valid xlsx (zip) file"

        # Verify structure with openpyxl
        try:
            from openpyxl import load_workbook
        except ImportError:
            pytest.skip("openpyxl not available")
        wb = load_workbook(io.BytesIO(r2.content), read_only=True)
        expected = {"Ringkasan per ACR", "Detail Asset", "Ringkasan Status"}
        assert expected.issubset(set(wb.sheetnames)), f"missing sheets, got {wb.sheetnames}"
        detail = wb["Detail Asset"]
        rows = list(detail.iter_rows(values_only=True))
        # first row is header
        data_rows = [r for r in rows[1:] if any(c is not None for c in r)]
        assert len(data_rows) >= 25, f"Detail Asset expected >=25 rows, got {len(data_rows)}"

    def test_bad_token(self):
        r = requests.get(f"{API}/rcg/reports/assets.xlsx", params={"token": "bad"}, timeout=15)
        assert r.status_code == 401, r.text


# ---------- Dashboard RCG by_status includes SOLD ----------
class TestDashboardSoldKey:
    def test_rcg_dash_has_sold(self, admin_token):
        r = requests.get(f"{API}/dashboard/rcg", headers=H(admin_token), timeout=15)
        assert r.status_code == 200
        d = r.json()
        assert "by_status" in d, f"missing by_status: {d}"
        assert "SOLD" in d["by_status"], f"SOLD key missing: {list(d['by_status'].keys())}"
