"""BSI ASSET DEAL — Iteration 5 tests.

Covers:
- GET /api/public/catalog?price_drop=true (and combined with provinsi=Bali)
- GET /api/rcg/reports/assets/link -> assets.xlsx: sheetnames, headers of
  'Detail Asset' & 'Minat Pembeli', row ordering & counts.
"""
import io
import os
import pytest
import requests

BASE_URL = os.environ["EXPO_PUBLIC_BACKEND_URL"].rstrip("/")
API = f"{BASE_URL}/api"


def _login(username, password):
    r = requests.post(f"{API}/auth/login",
                      json={"username": username, "password": password}, timeout=30)
    assert r.status_code == 200, f"login {username}: {r.status_code} {r.text}"
    return r.json()["access_token"]


def H(t):
    return {"Authorization": f"Bearer {t}"}


@pytest.fixture(scope="module")
def admin_token():
    return _login("admin", "Admin@2026")


@pytest.fixture(scope="module")
def ma_token():
    return _login("2190020284", "BSI@2026")


# ================================================================
# price_drop filter
# ================================================================
class TestPriceDropFilter:
    def test_catalog_no_flag_total_25(self):
        r = requests.get(f"{API}/public/catalog", params={"limit": 30}, timeout=20)
        assert r.status_code == 200
        j = r.json()
        assert j["total"] == 25, f"expected 25 total, got {j['total']}"

    def test_price_drop_true_all_have_drop(self):
        r = requests.get(f"{API}/public/catalog",
                         params={"price_drop": "true", "limit": 30}, timeout=20)
        assert r.status_code == 200
        j = r.json()
        items = j["items"]
        assert j["total"] == len(items), f"total {j['total']} != items {len(items)}"
        assert 3 <= len(items) <= 10, f"expected ~5 drops, got {len(items)}"
        for it in items:
            p = it.get("penurunan_persen") or 0
            assert p > 0, f"asset {it['id']} penurunan_persen not > 0: {p}"
            assert it.get("harga_sebelumnya") and it["harga_sebelumnya"] > it["harga_limit"]

    def test_price_drop_false_same_as_default(self):
        r_def = requests.get(f"{API}/public/catalog", params={"limit": 30}, timeout=20).json()
        r_false = requests.get(f"{API}/public/catalog",
                               params={"price_drop": "false", "limit": 30}, timeout=20).json()
        assert r_def["total"] == r_false["total"] == 25

    def test_price_drop_with_provinsi_bali(self):
        r = requests.get(f"{API}/public/catalog",
                         params={"price_drop": "true", "provinsi": "Bali", "limit": 30}, timeout=20)
        assert r.status_code == 200
        j = r.json()
        for it in j["items"]:
            assert it.get("provinsi") == "Bali", f"non-Bali leaked: {it.get('provinsi')}"
            assert (it.get("penurunan_persen") or 0) > 0
        # total consistent with items count in single page
        assert j["total"] == len(j["items"])


# ================================================================
# Excel report structure & headers
# ================================================================
class TestReportXlsx:
    @pytest.fixture(scope="class")
    def wb(self, admin_token):
        r = requests.get(f"{API}/rcg/reports/assets/link",
                         headers=H(admin_token), timeout=20)
        assert r.status_code == 200, r.text
        j = r.json()
        assert j.get("url"), j
        url = j["url"]
        # url comes back like "/api/rcg/reports/assets.xlsx?token=..."
        full = url if url.startswith("http") else f"{BASE_URL}{url}"
        rr = requests.get(full, timeout=30)
        assert rr.status_code == 200, f"xlsx GET failed: {rr.status_code}"
        assert "spreadsheet" in rr.headers.get("content-type", "").lower() or \
               "officedocument" in rr.headers.get("content-type", "").lower(), rr.headers
        from openpyxl import load_workbook
        return load_workbook(io.BytesIO(rr.content), read_only=False, data_only=True)

    def test_sheet_names(self, wb):
        assert wb.sheetnames == ["Ringkasan per ACR", "Detail Asset",
                                 "Minat Pembeli", "Ringkasan Status"], wb.sheetnames

    def test_detail_asset_header(self, wb):
        ws = wb["Detail Asset"]
        header = [c.value for c in ws[1]]
        # last 4 columns must be the interest ones
        assert header[-4:] == ["Dilihat", "Ketuk WhatsApp",
                               "Dilihat 7 Hari", "WA 7 Hari"], header[-6:]

    def test_minat_pembeli_header(self, wb):
        ws = wb["Minat Pembeli"]
        header = [c.value for c in ws[1]]
        expected = ["Peringkat", "Nomor Asset", "Judul", "Status", "ACR",
                    "PIC Marketing", "Harga Limit", "Dilihat", "Ketuk WhatsApp",
                    "Dilihat 7 Hari", "WA 7 Hari", "Konversi WA (%)"]
        assert header == expected, header

    def test_minat_pembeli_sorted_and_row_count(self, wb, admin_token):
        ws = wb["Minat Pembeli"]
        # Data rows exclude header
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert len(rows) > 0, "no data rows in Minat Pembeli"

        # Check sort: score = wa*5 + views (descending)
        scores = []
        for r in rows:
            # Peringkat, Nomor, Judul, Status, ACR, PIC, HargaLimit, Dilihat, WA, Dil7, WA7, Konv
            views = r[7] or 0
            wa = r[8] or 0
            scores.append(wa * 5 + views)
        for i in range(1, len(scores)):
            assert scores[i - 1] >= scores[i], \
                f"not desc sorted at row {i}: {scores[i-1]} < {scores[i]}"

        # row count == number of assets (all statuses that appear in reports)
        # Compare against Ringkasan Status total
        ws3 = wb["Ringkasan Status"]
        total_assets = 0
        for r in ws3.iter_rows(min_row=2, values_only=True):
            total_assets += (r[1] or 0)
        assert len(rows) == total_assets, \
            f"Minat Pembeli rows ({len(rows)}) != total assets in report ({total_assets})"
