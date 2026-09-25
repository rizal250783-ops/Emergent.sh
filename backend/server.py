from fastapi import FastAPI, APIRouter, Depends, HTTPException, UploadFile, File, Form, Query
from fastapi.responses import Response as FastResponse
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from starlette.concurrency import run_in_threadpool
from motor.motor_asyncio import AsyncIOMotorClient
import os, re, uuid, logging, jwt, bcrypt, requests, io
from pathlib import Path
from datetime import datetime, timezone, timedelta
from typing import Optional, List
from pydantic import BaseModel

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

from seed_data import ACR_ROWS, KPKNL_ROWS, CATEGORY_SEED

mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

JWT_SECRET = os.environ.get('JWT_SECRET', 'bsi-asset-deal-dev-secret')
JWT_ALG = 'HS256'
ACCESS_HOURS = 12
DEFAULT_PASSWORD = os.environ.get('DEFAULT_PASSWORD', 'BSI@2026')
ADMIN_PASSWORD = os.environ.get('ADMIN_PASSWORD', 'Admin@2026')

# ---------------- Object storage ----------------
STORAGE_BASE = (os.environ.get("INTEGRATION_PROXY_URL") or "").strip() or "https://integrations.emergentagent.com"
STORAGE_URL = STORAGE_BASE.rstrip("/") + "/objstore/api/v1/storage"
EMERGENT_KEY = os.environ.get("EMERGENT_LLM_KEY")
APP_NAME = "bsi-asset-deal"
_storage_key = None

def init_storage():
    global _storage_key
    if _storage_key:
        return _storage_key
    resp = requests.post(f"{STORAGE_URL}/init", json={"emergent_key": EMERGENT_KEY}, timeout=30)
    resp.raise_for_status()
    _storage_key = resp.json()["storage_key"]
    return _storage_key

def put_object(path: str, data: bytes, content_type: str) -> dict:
    key = init_storage()
    resp = requests.put(f"{STORAGE_URL}/objects/{path}",
        headers={"X-Storage-Key": key, "Content-Type": content_type}, data=data, timeout=120)
    resp.raise_for_status()
    return resp.json()

def get_object(path: str):
    global _storage_key
    key = init_storage()
    resp = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    if resp.status_code == 503:
        _storage_key = None
        key = init_storage()
        resp = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    resp.raise_for_status()
    return resp.content, resp.headers.get("Content-Type", "application/octet-stream")

# ---------------- helpers ----------------
def now_iso():
    return datetime.now(timezone.utc).isoformat()

def new_id():
    return str(uuid.uuid4())

def clean(doc):
    if doc is None:
        return None
    doc.pop('_id', None)
    return doc

def hash_pw(pw: str) -> str:
    return bcrypt.hashpw(pw.encode(), bcrypt.gensalt()).decode()

def verify_pw(pw: str, hashed: str) -> bool:
    try:
        return bcrypt.checkpw(pw.encode(), hashed.encode())
    except Exception:
        return False

def make_token(user):
    payload = {"sub": user["username"], "role": user["role"], "uid": user["id"],
               "exp": datetime.now(timezone.utc) + timedelta(hours=ACCESS_HOURS)}
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALG)

def acr_identifier(nama_acr: str) -> str:
    s = nama_acr.upper().replace("ACR", "", 1) if nama_acr.upper().startswith("ACR") else nama_acr.upper()
    return re.sub(r'[^A-Z0-9]', '', s)

def normalize_wa(hp: str) -> str:
    d = re.sub(r'\D', '', hp or '')
    if d.startswith('0'):
        d = '62' + d[1:]
    elif not d.startswith('62'):
        d = '62' + d
    return d

app = FastAPI()
api = APIRouter(prefix="/api")
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("bsi")

# ---------------- auth dependency ----------------
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
security = HTTPBearer(auto_error=False)

async def current_user(creds: Optional[HTTPAuthorizationCredentials] = Depends(security)):
    if not creds:
        raise HTTPException(401, "Token diperlukan")
    try:
        payload = jwt.decode(creds.credentials, JWT_SECRET, algorithms=[JWT_ALG])
    except jwt.PyJWTError:
        raise HTTPException(401, "Token tidak valid atau kadaluarsa")
    user = await db.users.find_one({"id": payload.get("uid"), "status": "active", "deleted_at": None})
    if not user:
        raise HTTPException(401, "Akun tidak aktif")
    return clean(user)

def require(*roles):
    async def dep(user=Depends(current_user)):
        if user["role"] not in roles:
            raise HTTPException(403, "Anda tidak memiliki akses ke data ini")
        return user
    return dep

# ---------------- audit / notif ----------------
async def audit(user, action, entity, entity_id, before=None, after=None, notes=None):
    await db.audit_logs.insert_one({
        "id": new_id(), "user": user.get("username") if user else "system",
        "role": user.get("role") if user else "system", "action": action,
        "entity": entity, "entity_id": entity_id, "data_before": before,
        "data_after": after, "notes": notes, "timestamp": now_iso()})

async def notify(user_id, ntype, title, message):
    if not user_id:
        return
    await db.notifications.insert_one({"id": new_id(), "user_id": user_id, "type": ntype,
        "title": title, "message": message, "is_read": False, "created_at": now_iso()})

async def approval_log(asset_id, action, before, after, reviewer, notes=None):
    await db.approval_logs.insert_one({"id": new_id(), "asset_id": asset_id, "action": action,
        "status_before": before, "status_after": after,
        "reviewer_id": reviewer.get("id") if reviewer else None,
        "reviewer_name": reviewer.get("nama") if reviewer else None,
        "reviewer_role": reviewer.get("role") if reviewer else None,
        "notes": notes, "timestamp": now_iso()})

# =================== SEED ===================
@app.on_event("startup")
async def startup():
    await db.users.create_index("username", unique=True)
    try:
        init_storage()
    except Exception as e:
        logger.warning(f"storage init deferred: {e}")

    if await db.master_acr.count_documents({}) == 0:
        await seed_all()

async def seed_all():
    logger.info("Seeding master data...")
    # ACR + ACRM + MA
    acr_map = {}  # nama -> id
    acrm_first = {}  # nama_acr -> acrm row (first)
    for row in ACR_ROWS:
        nama_acr, nip_acrm, nama_acrm, nama_ma, nip_ma, hp_ma, placeholder = row
        if nama_acr not in acr_map:
            aid = new_id()
            acr_map[nama_acr] = aid
            await db.master_acr.insert_one({"id": aid, "nama_acr": nama_acr,
                "identifier": acr_identifier(nama_acr), "status": "active",
                "created_at": now_iso(), "updated_at": now_iso(), "deleted_at": None})
        if nama_acr not in acrm_first:
            acrm_first[nama_acr] = (nip_acrm, nama_acrm, placeholder)

    # ACRM (one per ACR)
    acrm_map = {}  # nama_acr -> acrm id
    for nama_acr, aid in acr_map.items():
        nip_acrm, nama_acrm, ph = acrm_first[nama_acr]
        acid = new_id()
        acrm_map[nama_acr] = acid
        flag = "PERLU_KONFIRMASI_DATA" if (ph or nama_acrm == "NOPE") else "OK"
        await db.master_acrm.insert_one({"id": acid, "id_acr": aid, "nama_acrm": nama_acrm,
            "nip_acrm": nip_acrm, "nomor_hp": "", "status": "active", "data_flag": flag,
            "created_at": now_iso(), "updated_at": now_iso(), "deleted_at": None})
        await create_user(acid, nama_acrm, nip_acrm, "acrm", flag)

    # Marketing Asset
    for row in ACR_ROWS:
        nama_acr, nip_acrm, nama_acrm, nama_ma, nip_ma, hp_ma, placeholder = row
        aid = acr_map[nama_acr]
        mid = new_id()
        flag = "PERLU_KONFIRMASI_DATA" if (placeholder or nama_ma == "NOPE") else "OK"
        await db.master_marketing_asset.insert_one({"id": mid, "id_acr": aid,
            "nama_marketing_asset": nama_ma, "nip": nip_ma, "nomor_hp": hp_ma,
            "status": "active", "data_flag": flag,
            "created_at": now_iso(), "updated_at": now_iso(), "deleted_at": None})
        await create_user(mid, nama_ma, nip_ma, "marketing_asset", flag)

    # KPKNL
    for nama, alamat in KPKNL_ROWS:
        prov = None
        m = re.search(r'Provinsi ([^,]+)', alamat)
        if m:
            prov = m.group(1).strip()
        await db.master_kpknl.insert_one({"id": new_id(), "nama_kpknl": nama,
            "alamat_kpknl": alamat, "provinsi": prov, "status": "active",
            "created_at": now_iso(), "updated_at": now_iso(), "deleted_at": None})

    # Categories
    for parent, subs in CATEGORY_SEED.items():
        pid = new_id()
        await db.master_asset_category.insert_one({"id": pid, "parent_category_id": None,
            "nama_category": parent, "status": "active", "created_at": now_iso(),
            "updated_at": now_iso(), "deleted_at": None})
        for s in subs:
            await db.master_asset_category.insert_one({"id": new_id(), "parent_category_id": pid,
                "nama_category": s, "status": "active", "created_at": now_iso(),
                "updated_at": now_iso(), "deleted_at": None})

    # Admin RCG users
    for uname, nama in [("admin", "Admin RCG Pusat"), ("admin2", "Admin RCG Pusat 2")]:
        if not await db.users.find_one({"username": uname}):
            await db.users.insert_one({"id": new_id(), "username": uname,
                "password_hash": hash_pw(ADMIN_PASSWORD), "role": "admin_rcg", "ref_id": None,
                "nama": nama, "status": "active", "data_flag": "OK", "last_login_at": None,
                "force_password_change": False, "created_at": now_iso(),
                "updated_at": now_iso(), "deleted_at": None})
    logger.info("Seeding done.")

async def create_user(ref_id, nama, nip, role, flag):
    uname = (nip or "").strip().lower()
    if not uname or await db.users.find_one({"username": uname}):
        # dummy placeholder nips may collide; make unique
        uname = f"{uname}-{ref_id[:6]}" if uname else ref_id[:8]
    await db.users.insert_one({"id": new_id(), "username": uname,
        "password_hash": hash_pw(DEFAULT_PASSWORD), "role": role, "ref_id": ref_id,
        "nama": nama, "status": "active", "data_flag": flag, "last_login_at": None,
        "force_password_change": True, "created_at": now_iso(),
        "updated_at": now_iso(), "deleted_at": None})

# =================== AUTH ===================
class LoginIn(BaseModel):
    username: str
    password: str

class ChangePwIn(BaseModel):
    old_password: str
    new_password: str

async def enrich_user(user):
    """attach acr/acrm context for internal roles"""
    out = {"id": user["id"], "username": user["username"], "role": user["role"],
           "nama": user.get("nama"), "data_flag": user.get("data_flag"),
           "force_password_change": user.get("force_password_change", False)}
    if user["role"] == "marketing_asset" and user.get("ref_id"):
        ma = await db.master_marketing_asset.find_one({"id": user["ref_id"]})
        if ma:
            acr = await db.master_acr.find_one({"id": ma["id_acr"]})
            acrm = await db.master_acrm.find_one({"id_acr": ma["id_acr"], "status": "active", "deleted_at": None})
            out["ma"] = {"id": ma["id"], "nama": ma["nama_marketing_asset"], "nip": ma["nip"], "hp": ma["nomor_hp"]}
            out["acr"] = {"id": acr["id"], "nama": acr["nama_acr"], "identifier": acr["identifier"]} if acr else None
            out["acrm"] = {"id": acrm["id"], "nama": acrm["nama_acrm"]} if acrm else None
    elif user["role"] == "acrm" and user.get("ref_id"):
        acrm = await db.master_acrm.find_one({"id": user["ref_id"]})
        if acrm:
            acr = await db.master_acr.find_one({"id": acrm["id_acr"]})
            out["acrm"] = {"id": acrm["id"], "nama": acrm["nama_acrm"], "nip": acrm["nip_acrm"]}
            out["acr"] = {"id": acr["id"], "nama": acr["nama_acr"]} if acr else None
    return out

@api.post("/auth/login")
async def login(body: LoginIn):
    uname = body.username.strip().lower()
    user = await db.users.find_one({"username": uname, "deleted_at": None})
    if not user or not verify_pw(body.password, user["password_hash"]):
        raise HTTPException(401, "NIP/username atau password salah")
    if user["status"] != "active":
        raise HTTPException(403, "Akun Anda tidak aktif. Hubungi Admin RCG.")
    await db.users.update_one({"id": user["id"]}, {"$set": {"last_login_at": now_iso()}})
    await audit(user, "LOGIN", "users", user["id"])
    return {"access_token": make_token(user), "user": await enrich_user(user)}

@api.get("/auth/me")
async def me(user=Depends(current_user)):
    full = await db.users.find_one({"id": user["id"]})
    return await enrich_user(full)

@api.post("/auth/change-password")
async def change_pw(body: ChangePwIn, user=Depends(current_user)):
    full = await db.users.find_one({"id": user["id"]})
    if not verify_pw(body.old_password, full["password_hash"]):
        raise HTTPException(400, "Password lama salah")
    if len(body.new_password) < 6:
        raise HTTPException(400, "Password baru minimal 6 karakter")
    await db.users.update_one({"id": user["id"]}, {"$set": {
        "password_hash": hash_pw(body.new_password), "force_password_change": False, "updated_at": now_iso()}})
    await audit(user, "CHANGE_PASSWORD", "users", user["id"])
    return {"ok": True}

# =================== MASTER (read) ===================
@api.get("/master/categories")
async def categories(active_only: bool = True):
    q = {"deleted_at": None}
    if active_only:
        q["status"] = "active"
    cats = [clean(c) async for c in db.master_asset_category.find(q)]
    parents = [c for c in cats if not c["parent_category_id"]]
    for p in parents:
        p["subcategories"] = [c for c in cats if c["parent_category_id"] == p["id"]]
    return parents

@api.get("/master/kpknl")
async def kpknl_list(user=Depends(current_user)):
    return [clean(k) async for k in db.master_kpknl.find({"status": "active", "deleted_at": None}).sort("nama_kpknl", 1)]

# =================== ASSET helpers ===================
ASSET_PUBLIC_STATUSES = ["PUBLISHED", "UPDATE_PENDING_ACRM", "UPDATE_PENDING_RCG", "SOLD"]

async def gen_asset_number(acr):
    year = datetime.now().year
    ident = acr["identifier"]
    seq = await db.counters.find_one_and_update(
        {"_id": f"asset-{year}-{ident}"}, {"$inc": {"seq": 1}},
        upsert=True, return_document=True)
    n = seq["seq"]
    return f"BLC-{year}-{ident}-{str(n).zfill(6)}"

async def asset_public_view(a):
    imgs = [clean(i) async for i in db.asset_images.find({"id_asset": a["id"], "deleted_at": None})]
    cat = await db.master_asset_category.find_one({"id": a.get("id_category")})
    sub = await db.master_asset_category.find_one({"id": a.get("id_subcategory")}) if a.get("id_subcategory") else None
    sched = a.get("schedule") or {}
    kpknl = None
    if sched.get("id_kpknl"):
        k = await db.master_kpknl.find_one({"id": sched["id_kpknl"]})
        kpknl = {"nama": k["nama_kpknl"], "alamat": k["alamat_kpknl"]} if k else None
    return {
        "id": a["id"], "nomor_asset": a["nomor_asset"], "judul_asset": a["judul_asset"],
        "deskripsi": a.get("deskripsi"), "kategori": cat["nama_category"] if cat else None,
        "subkategori": sub["nama_category"] if sub else None,
        "alamat": a.get("alamat"), "provinsi": a.get("provinsi"),
        "kabupaten_kota": a.get("kabupaten_kota"), "kecamatan": a.get("kecamatan"),
        "wilayah_level_4": a.get("wilayah_level_4"), "tipe_wilayah": a.get("tipe_wilayah"),
        "latitude": a.get("latitude"), "longitude": a.get("longitude"),
        "luas_tanah": a.get("luas_tanah"), "luas_bangunan": a.get("luas_bangunan"),
        "kondisi_asset": a.get("kondisi_asset"), "harga_limit": a.get("harga_limit"),
        "nilai_appraisal": a.get("nilai_appraisal"),
        "extra": a.get("extra") or {},
        "images": [i["url"] for i in imgs],
        "status": a.get("status"),
        "is_sold": a.get("status") == "SOLD",
        "sold_at": a.get("sold_at"),
        "has_schedule": bool(sched.get("tanggal_lelang")),
        "tanggal_lelang": sched.get("tanggal_lelang"),
        "kpknl": kpknl,
        "pic_nama": a.get("pic_nama"),
        # SOLD: hide PIC contact so marketing stops receiving inquiries
        "pic_wa": normalize_wa(a.get("pic_hp")) if (a.get("pic_hp") and a.get("status") != "SOLD") else None,
    }

# =================== PUBLIC CATALOG ===================
@api.get("/public/filters")
async def public_filters():
    match = {"status": {"$in": ASSET_PUBLIC_STATUSES}, "public_ready": True, "deleted_at": None}
    provinsi = await db.assets.distinct("provinsi", match)
    cats = [clean(c) async for c in db.master_asset_category.find({"parent_category_id": None, "status": "active", "deleted_at": None})]
    return {"provinsi": sorted([p for p in provinsi if p]), "categories": cats}

@api.get("/public/catalog")
async def public_catalog(keyword: Optional[str] = None, category_id: Optional[str] = None,
    subcategory_id: Optional[str] = None, provinsi: Optional[str] = None,
    kabupaten_kota: Optional[str] = None, kecamatan: Optional[str] = None,
    wilayah_level_4: Optional[str] = None, sort: str = "newest",
    page: int = 1, limit: int = 20):
    q = {"status": {"$in": ASSET_PUBLIC_STATUSES}, "public_ready": True, "deleted_at": None}
    if category_id: q["id_category"] = category_id
    if subcategory_id: q["id_subcategory"] = subcategory_id
    if provinsi: q["provinsi"] = provinsi
    if kabupaten_kota: q["kabupaten_kota"] = kabupaten_kota
    if kecamatan: q["kecamatan"] = kecamatan
    if wilayah_level_4: q["wilayah_level_4"] = wilayah_level_4
    if keyword:
        rx = {"$regex": re.escape(keyword), "$options": "i"}
        q["$or"] = [{"judul_asset": rx}, {"alamat": rx}, {"nomor_asset": rx}, {"provinsi": rx}, {"kabupaten_kota": rx}]
    sort_map = {"newest": ("published_at", -1), "price_low": ("harga_limit", 1),
                "price_high": ("harga_limit", -1), "auction": ("schedule.tanggal_lelang", 1)}
    sf, sd = sort_map.get(sort, ("published_at", -1))
    total = await db.assets.count_documents(q)
    # SOLD assets always sink to the end so buyers see available assets first
    pipeline = [{"$match": q},
        {"$addFields": {"_sold_rank": {"$cond": [{"$eq": ["$status", "SOLD"]}, 1, 0]}}},
        {"$sort": {"_sold_rank": 1, sf: sd, "_id": 1}},
        {"$skip": (page - 1) * limit}, {"$limit": limit}]
    items = [await asset_public_view(clean(a)) async for a in db.assets.aggregate(pipeline)]
    return {"total": total, "page": page, "limit": limit, "items": items}

@api.get("/public/catalog/batch")
async def public_batch(ids: str):
    """Fetch several published assets by id (favorites stored on-device)."""
    id_list = [i for i in ids.split(",") if i][:50]
    q = {"id": {"$in": id_list}, "status": {"$in": ASSET_PUBLIC_STATUSES}, "public_ready": True, "deleted_at": None}
    found = {a["id"]: a async for a in db.assets.find(q)}
    return [await asset_public_view(clean(found[i])) for i in id_list if i in found]

@api.get("/public/catalog/{asset_id}/similar")
async def public_similar(asset_id: str, limit: int = 6):
    """Similar assets: same location / category / price band, scored and ranked."""
    a = await db.assets.find_one({"id": asset_id, "deleted_at": None})
    if not a:
        raise HTTPException(404, "Asset tidak ditemukan")
    price = a.get("harga_limit") or 0
    q = {"id": {"$ne": asset_id}, "status": {"$in": ["PUBLISHED", "UPDATE_PENDING_ACRM", "UPDATE_PENDING_RCG"]},
         "public_ready": True, "deleted_at": None,
         "$or": [{"provinsi": a.get("provinsi")}, {"id_category": a.get("id_category")},
                 {"harga_limit": {"$gte": price * 0.6, "$lte": price * 1.4}}]}
    scored = []
    async for c in db.assets.find(q):
        s = 0
        if c.get("kabupaten_kota") == a.get("kabupaten_kota"): s += 4
        if c.get("provinsi") == a.get("provinsi"): s += 3
        if c.get("id_subcategory") == a.get("id_subcategory"): s += 3
        elif c.get("id_category") == a.get("id_category"): s += 2
        if price and c.get("harga_limit") and price * 0.6 <= c["harga_limit"] <= price * 1.4: s += 2
        scored.append((s, c))
    scored.sort(key=lambda x: (-x[0], x[1].get("published_at") or ""), reverse=False)
    return [await asset_public_view(clean(c)) for _, c in scored[:limit]]

@api.get("/public/catalog/{asset_id}")
async def public_detail(asset_id: str):
    a = await db.assets.find_one({"id": asset_id, "status": {"$in": ASSET_PUBLIC_STATUSES},
        "public_ready": True, "deleted_at": None})
    if not a:
        raise HTTPException(404, "Asset tidak ditemukan")
    return await asset_public_view(clean(a))

@api.get("/public/locations")
async def public_locations(provinsi: Optional[str] = None, kabupaten_kota: Optional[str] = None,
    kecamatan: Optional[str] = None):
    """Cascading location options derived from published assets only (no empty results)."""
    match = {"status": {"$in": ASSET_PUBLIC_STATUSES}, "public_ready": True, "deleted_at": None}
    if not provinsi:
        field = "provinsi"
    elif not kabupaten_kota:
        match["provinsi"] = provinsi; field = "kabupaten_kota"
    elif not kecamatan:
        match["provinsi"] = provinsi; match["kabupaten_kota"] = kabupaten_kota; field = "kecamatan"
    else:
        match["provinsi"] = provinsi; match["kabupaten_kota"] = kabupaten_kota
        match["kecamatan"] = kecamatan; field = "wilayah_level_4"
    values = await db.assets.distinct(field, match)
    return {"level": field, "options": sorted([v for v in values if v])}

# =================== WILAYAH (administrative reference, cached proxy) ===================
WILAYAH_SRC = "https://www.emsifa.com/api-wilayah-indonesia/api"
_KEEP_UPPER = {"DKI", "DI", "KEP.", "KEP"}

def title_wilayah(name: str) -> str:
    out = []
    for w in (name or "").split():
        out.append(w if w in _KEEP_UPPER else w.capitalize())
    return " ".join(out)

def fetch_wilayah(kind: str, parent_id: Optional[str]):
    url = f"{WILAYAH_SRC}/{kind}.json" if kind == "provinces" else f"{WILAYAH_SRC}/{kind}/{parent_id}.json"
    r = requests.get(url, timeout=20)
    r.raise_for_status()
    return r.json()

async def wilayah_list(kind: str, parent_id: Optional[str] = None):
    key = f"{kind}:{parent_id or 'root'}"
    cached = await db.wilayah_cache.find_one({"_id": key})
    if cached:
        return cached["items"]
    try:
        raw = await run_in_threadpool(fetch_wilayah, kind, parent_id)
    except Exception as e:
        logger.warning(f"wilayah fetch failed {key}: {e}")
        raise HTTPException(503, "Data wilayah sementara tidak tersedia. Coba lagi.")
    items = [{"id": x["id"], "name": title_wilayah(x["name"])} for x in raw]
    await db.wilayah_cache.update_one({"_id": key}, {"$set": {"items": items, "cached_at": now_iso()}}, upsert=True)
    return items

@api.get("/wilayah/provinces")
async def wilayah_provinces():
    return await wilayah_list("provinces")

@api.get("/wilayah/regencies/{province_id}")
async def wilayah_regencies(province_id: str):
    return await wilayah_list("regencies", province_id)

@api.get("/wilayah/districts/{regency_id}")
async def wilayah_districts(regency_id: str):
    return await wilayah_list("districts", regency_id)

@api.get("/wilayah/villages/{district_id}")
async def wilayah_villages(district_id: str):
    return await wilayah_list("villages", district_id)

# =================== FILE ===================
@api.get("/files/private/{doc_id}")
async def serve_private_file(doc_id: str, token: str = Query(...)):
    """Signed-URL access for private legal documents (short-lived token)."""
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALG])
    except jwt.PyJWTError:
        raise HTTPException(401, "Tautan tidak valid atau sudah kadaluarsa")
    if payload.get("typ") != "doc" or payload.get("doc") != doc_id:
        raise HTTPException(401, "Tautan tidak valid")
    d = await db.asset_documents.find_one({"id": doc_id, "deleted_at": None})
    if not d:
        raise HTTPException(404, "Dokumen tidak ditemukan")
    try:
        content, ctype = await run_in_threadpool(get_object, d["storage_path"])
    except Exception:
        raise HTTPException(404, "File tidak ditemukan")
    disp = "inline" if ctype.startswith("image/") or ctype == "application/pdf" else "attachment"
    return FastResponse(content=content, media_type=ctype,
        headers={"Content-Disposition": f'{disp}; filename="{d["nama_file"]}"'})

@api.get("/files/{path:path}")
async def serve_file(path: str):
    if "/private/" in f"/{path}":
        raise HTTPException(403, "Dokumen ini bersifat privat")
    try:
        content, ctype = await run_in_threadpool(get_object, path)
    except Exception:
        raise HTTPException(404, "File tidak ditemukan")
    return FastResponse(content=content, media_type=ctype)

# =================== MARKETING ASSET ===================
class AssetIn(BaseModel):
    judul_asset: str
    deskripsi: str
    id_category: str
    id_subcategory: Optional[str] = None
    alamat: str
    provinsi: str
    kabupaten_kota: str
    kecamatan: str
    wilayah_level_4: str
    tipe_wilayah: str = "Kelurahan"
    latitude: Optional[float] = None
    longitude: Optional[float] = None
    luas_tanah: Optional[float] = None
    luas_bangunan: Optional[float] = None
    kondisi_asset: Optional[str] = None
    nilai_appraisal: Optional[float] = None
    harga_limit: Optional[float] = None
    extra: Optional[dict] = None

async def ma_context(user):
    full = await db.users.find_one({"id": user["id"]})
    ma = await db.master_marketing_asset.find_one({"id": full["ref_id"]})
    if not ma:
        raise HTTPException(400, "Data Marketing Asset tidak ditemukan")
    acr = await db.master_acr.find_one({"id": ma["id_acr"]})
    acrm = await db.master_acrm.find_one({"id_acr": ma["id_acr"], "status": "active", "deleted_at": None})
    return ma, acr, acrm

@api.get("/marketing/context")
async def marketing_context(user=Depends(require("marketing_asset"))):
    ma, acr, acrm = await ma_context(user)
    return {"nama_acr": acr["nama_acr"] if acr else None,
            "nama_acrm": acrm["nama_acrm"] if acrm else None,
            "nama_marketing_asset": ma["nama_marketing_asset"],
            "nomor_hp": ma["nomor_hp"]}

def validate_asset_numbers(body: AssetIn):
    for f in ["luas_tanah", "luas_bangunan", "nilai_appraisal", "harga_limit"]:
        v = getattr(body, f)
        if v is not None and v <= 0:
            raise HTTPException(400, f"{f} harus lebih besar dari 0")

@api.post("/assets")
async def create_asset(body: AssetIn, user=Depends(require("marketing_asset"))):
    validate_asset_numbers(body)
    ma, acr, acrm = await ma_context(user)
    num = await gen_asset_number(acr)
    a = {"id": new_id(), "nomor_asset": num, "id_marketing_asset": ma["id"],
         "id_acr": acr["id"], "id_acrm": acrm["id"] if acrm else None,
         "acr_nama": acr["nama_acr"], "acrm_nama": acrm["nama_acrm"] if acrm else None,
         "pic_nama": ma["nama_marketing_asset"], "pic_hp": ma["nomor_hp"],
         "status": "DRAFT", "public_ready": False, "current_version_no": 1,
         "schedule": {}, "correction_notes": None, "published_at": None,
         "created_at": now_iso(), "updated_at": now_iso(), "deleted_at": None,
         **body.model_dump()}
    await db.assets.insert_one(a)
    await audit(user, "CREATE", "assets", a["id"], after={"nomor_asset": num})
    await notify(user["id"], "asset", "Asset dibuat", f"Asset {num} berhasil disimpan sebagai draft.")
    return clean(a)

@api.put("/assets/{asset_id}")
async def update_asset(asset_id: str, body: AssetIn, user=Depends(require("marketing_asset"))):
    validate_asset_numbers(body)
    a = await db.assets.find_one({"id": asset_id, "deleted_at": None})
    if not a:
        raise HTTPException(404, "Asset tidak ditemukan")
    ma, acr, acrm = await ma_context(user)
    if a["id_marketing_asset"] != ma["id"]:
        raise HTTPException(403, "Anda tidak memiliki akses ke asset ini")
    if a["status"] not in ["DRAFT", "RETURN_TO_MARKETING", "RETURN_FROM_RCG", "PUBLISHED", "SOLD"]:
        raise HTTPException(400, "Asset sedang dalam proses review, tidak dapat diedit")
    before = {k: a.get(k) for k in body.model_dump()}
    await db.assets.update_one({"id": asset_id}, {"$set": {**body.model_dump(), "updated_at": now_iso()}})
    await audit(user, "UPDATE", "assets", asset_id, before=before, after=body.model_dump())
    return clean(await db.assets.find_one({"id": asset_id}))

@api.post("/assets/{asset_id}/schedule")
async def set_schedule(asset_id: str, tanggal_lelang: str = Form(...), id_kpknl: str = Form(...),
    user=Depends(require("marketing_asset"))):
    a = await db.assets.find_one({"id": asset_id, "deleted_at": None})
    if not a:
        raise HTTPException(404, "Asset tidak ditemukan")
    ma, acr, acrm = await ma_context(user)
    if a["id_marketing_asset"] != ma["id"]:
        raise HTTPException(403, "Anda tidak memiliki akses ke asset ini")
    try:
        datetime.fromisoformat(tanggal_lelang)
    except Exception:
        raise HTTPException(400, "Format tanggal lelang tidak valid")
    k = await db.master_kpknl.find_one({"id": id_kpknl, "status": "active", "deleted_at": None})
    if not k:
        raise HTTPException(400, "KPKNL tidak valid / tidak aktif")
    await db.assets.update_one({"id": asset_id}, {"$set": {
        "schedule": {"tanggal_lelang": tanggal_lelang, "id_kpknl": id_kpknl, "status_jadwal": "SUDAH ADA JADWAL"},
        "updated_at": now_iso()}})
    await audit(user, "SET_SCHEDULE", "assets", asset_id, after={"tanggal_lelang": tanggal_lelang, "id_kpknl": id_kpknl})
    return clean(await db.assets.find_one({"id": asset_id}))

@api.post("/assets/{asset_id}/submit")
async def submit_asset(asset_id: str, user=Depends(require("marketing_asset"))):
    a = await db.assets.find_one({"id": asset_id, "deleted_at": None})
    if not a:
        raise HTTPException(404, "Asset tidak ditemukan")
    ma, acr, acrm = await ma_context(user)
    if a["id_marketing_asset"] != ma["id"]:
        raise HTTPException(403, "Anda tidak memiliki akses ke asset ini")
    if not acrm:
        raise HTTPException(400, "Belum ada ACRM aktif untuk ACR Anda")
    before = a["status"]
    is_update = a["status"] in ["PUBLISHED", "SOLD"]
    new_status = "UPDATE_PENDING_ACRM" if is_update else "WAITING_ACRM_REVIEW"
    if a["status"] not in ["DRAFT", "RETURN_TO_MARKETING", "RETURN_FROM_RCG", "PUBLISHED", "SOLD"]:
        raise HTTPException(400, "Status asset tidak memungkinkan submit")
    await db.assets.update_one({"id": asset_id}, {"$set": {"status": new_status,
        "correction_notes": None, "updated_at": now_iso()}})
    action = "UPDATE_SUBMIT" if is_update else ("RESUBMIT" if before in ["RETURN_TO_MARKETING", "RETURN_FROM_RCG"] else "SUBMIT")
    await approval_log(asset_id, action, before, new_status, {"id": user["id"], "nama": ma["nama_marketing_asset"], "role": "marketing_asset"})
    await audit(user, action, "assets", asset_id)
    acrm_user = await db.users.find_one({"ref_id": acrm["id"], "role": "acrm"})
    if acrm_user:
        await notify(acrm_user["id"], "review", "Asset menunggu review", f"Asset {a['nomor_asset']} menunggu review Anda.")
    return {"ok": True, "status": new_status}

@api.get("/assets/mine")
async def my_assets(status: Optional[str] = None, user=Depends(require("marketing_asset"))):
    ma, acr, acrm = await ma_context(user)
    q = {"id_marketing_asset": ma["id"], "deleted_at": None}
    if status:
        q["status"] = status
    items = []
    async for a in db.assets.find(q).sort("updated_at", -1):
        items.append(await asset_internal_view(clean(a)))
    return items

async def asset_internal_view(a):
    v = await asset_public_view(a)
    v["correction_notes"] = a.get("correction_notes")
    v["acr_nama"] = a.get("acr_nama")
    v["acrm_nama"] = a.get("acrm_nama")
    v["created_at"] = a.get("created_at")
    v["updated_at"] = a.get("updated_at")
    v["id_category"] = a.get("id_category")
    v["id_subcategory"] = a.get("id_subcategory")
    v["schedule_kpknl_id"] = (a.get("schedule") or {}).get("id_kpknl")
    return v

@api.get("/assets/{asset_id}")
async def get_asset_internal(asset_id: str, user=Depends(current_user)):
    a = await assert_asset_access(asset_id, user)
    v = await asset_internal_view(clean(a))
    logs = [clean(l) async for l in db.approval_logs.find({"asset_id": asset_id}).sort("timestamp", 1)]
    v["approval_history"] = logs
    v["documents"] = [doc_view(d) async for d in db.asset_documents.find({"id_asset": asset_id, "deleted_at": None}).sort("created_at", 1)]
    return v

async def assert_asset_access(asset_id: str, user):
    """Internal access control: MA owner, ACRM of same ACR, or Admin RCG."""
    a = await db.assets.find_one({"id": asset_id, "deleted_at": None})
    if not a:
        raise HTTPException(404, "Asset tidak ditemukan")
    if user["role"] == "marketing_asset":
        ma = await db.master_marketing_asset.find_one({"id": (await db.users.find_one({"id": user["id"]}))["ref_id"]})
        if not ma or a["id_marketing_asset"] != ma["id"]:
            raise HTTPException(403, "Anda tidak memiliki akses ke asset ini")
    elif user["role"] == "acrm":
        acrm = await db.master_acrm.find_one({"id": (await db.users.find_one({"id": user["id"]}))["ref_id"]})
        if not acrm or a["id_acr"] != acrm["id_acr"]:
            raise HTTPException(403, "Anda tidak memiliki akses ke asset ini")
    return a

# ---------- Private legal documents (internal only) ----------
DOC_TYPES = {"application/pdf": "pdf", "image/jpeg": "jpg", "image/jpg": "jpg", "image/png": "png"}
DOC_EDITABLE_STATUSES = ["DRAFT", "RETURN_TO_MARKETING", "RETURN_FROM_RCG", "PUBLISHED", "SOLD"]

def doc_view(d):
    d = clean(d)
    d.pop("storage_path", None)
    return d

@api.post("/assets/{asset_id}/documents")
async def upload_document(asset_id: str, file: UploadFile = File(...), jenis: str = Form("Sertifikat"),
    user=Depends(require("marketing_asset"))):
    a = await assert_asset_access(asset_id, user)
    if file.content_type not in DOC_TYPES:
        raise HTTPException(400, "Format tidak didukung. Gunakan PDF, JPG atau PNG.")
    data = await file.read()
    if len(data) > 20 * 1024 * 1024:
        raise HTTPException(400, "Ukuran dokumen maksimal 20MB")
    ext = DOC_TYPES[file.content_type]
    path = f"{APP_NAME}/private/{asset_id}/{new_id()}.{ext}"
    await run_in_threadpool(put_object, path, data, file.content_type)
    d = {"id": new_id(), "id_asset": asset_id, "nama_file": file.filename or f"dokumen.{ext}",
         "jenis_dokumen": jenis, "content_type": file.content_type, "size": len(data),
         "storage_path": path, "is_public": False, "uploaded_by": user["id"],
         "uploaded_by_nama": a.get("pic_nama"), "created_at": now_iso(), "deleted_at": None}
    await db.asset_documents.insert_one(d)
    await audit(user, "UPLOAD_DOCUMENT", "asset_documents", d["id"], after={"asset": asset_id, "jenis": jenis, "nama_file": d["nama_file"]})
    return doc_view(d)

@api.get("/assets/{asset_id}/documents")
async def list_documents(asset_id: str, user=Depends(current_user)):
    await assert_asset_access(asset_id, user)
    return [doc_view(d) async for d in db.asset_documents.find({"id_asset": asset_id, "deleted_at": None}).sort("created_at", 1)]

@api.get("/assets/{asset_id}/documents/{doc_id}/link")
async def document_link(asset_id: str, doc_id: str, user=Depends(current_user)):
    await assert_asset_access(asset_id, user)
    d = await db.asset_documents.find_one({"id": doc_id, "id_asset": asset_id, "deleted_at": None})
    if not d:
        raise HTTPException(404, "Dokumen tidak ditemukan")
    token = jwt.encode({"typ": "doc", "doc": doc_id, "uid": user["id"],
        "exp": datetime.now(timezone.utc) + timedelta(minutes=10)}, JWT_SECRET, algorithm=JWT_ALG)
    await audit(user, "VIEW_DOCUMENT", "asset_documents", doc_id)
    return {"url": f"/api/files/private/{doc_id}?token={token}", "expires_in": 600}

@api.delete("/assets/{asset_id}/documents/{doc_id}")
async def delete_document(asset_id: str, doc_id: str, user=Depends(require("marketing_asset"))):
    a = await assert_asset_access(asset_id, user)
    if a["status"] not in DOC_EDITABLE_STATUSES:
        raise HTTPException(400, "Asset sedang dalam proses review, dokumen tidak dapat dihapus")
    r = await db.asset_documents.update_one({"id": doc_id, "id_asset": asset_id, "deleted_at": None},
        {"$set": {"deleted_at": now_iso(), "deleted_by": user["id"]}})
    if r.matched_count == 0:
        raise HTTPException(404, "Dokumen tidak ditemukan")
    await audit(user, "DELETE_DOCUMENT", "asset_documents", doc_id)
    return {"ok": True}

@api.post("/assets/{asset_id}/images")
async def upload_image(asset_id: str, file: UploadFile = File(...), jenis: str = Form("tambahan"),
    user=Depends(require("marketing_asset"))):
    a = await db.assets.find_one({"id": asset_id, "deleted_at": None})
    if not a:
        raise HTTPException(404, "Asset tidak ditemukan")
    ma = await db.master_marketing_asset.find_one({"id": (await db.users.find_one({"id": user["id"]}))["ref_id"]})
    if not ma or a["id_marketing_asset"] != ma["id"]:
        raise HTTPException(403, "Anda tidak memiliki akses ke asset ini")
    if file.content_type not in ["image/jpeg", "image/jpg", "image/png"]:
        raise HTTPException(400, "Format tidak didukung. Gunakan JPG atau PNG.")
    data = await file.read()
    if len(data) > 10 * 1024 * 1024:
        raise HTTPException(400, "Ukuran file maksimal 10MB")
    ext = "png" if file.content_type == "image/png" else "jpg"
    path = f"{APP_NAME}/uploads/{user['id']}/{new_id()}.{ext}"
    await run_in_threadpool(put_object, path, data, file.content_type)
    url = f"/api/files/{path}"
    img = {"id": new_id(), "id_asset": asset_id, "url": url, "storage_path": path,
           "jenis_foto": jenis, "is_public": True, "uploaded_by": user["id"],
           "created_at": now_iso(), "deleted_at": None}
    await db.asset_images.insert_one(img)
    return clean(img)

# =================== ACRM ===================
class ReviewIn(BaseModel):
    notes: Optional[str] = None

async def acrm_of(user):
    full = await db.users.find_one({"id": user["id"]})
    acrm = await db.master_acrm.find_one({"id": full["ref_id"]})
    if not acrm:
        raise HTTPException(400, "Data ACRM tidak ditemukan")
    return acrm

@api.get("/acrm/pending")
async def acrm_pending(user=Depends(require("acrm"))):
    acrm = await acrm_of(user)
    q = {"id_acr": acrm["id_acr"], "status": {"$in": ["WAITING_ACRM_REVIEW", "UPDATE_PENDING_ACRM"]}, "deleted_at": None}
    return [await asset_internal_view(clean(a)) async for a in db.assets.find(q).sort("updated_at", 1)]

@api.post("/acrm/assets/{asset_id}/approve")
async def acrm_approve(asset_id: str, user=Depends(require("acrm"))):
    acrm = await acrm_of(user)
    a = await db.assets.find_one({"id": asset_id, "deleted_at": None})
    if not a or a["id_acr"] != acrm["id_acr"]:
        raise HTTPException(403, "Anda tidak memiliki akses ke asset ini")
    if a["status"] not in ["WAITING_ACRM_REVIEW", "UPDATE_PENDING_ACRM"]:
        raise HTTPException(409, "Status asset sudah berubah, muat ulang data.")
    is_update = a["status"] == "UPDATE_PENDING_ACRM"
    new_status = "UPDATE_PENDING_RCG" if is_update else "WAITING_RCG_APPROVAL"
    await db.assets.update_one({"id": asset_id}, {"$set": {"status": new_status, "updated_at": now_iso()}})
    await approval_log(asset_id, "APPROVE", a["status"], new_status, {"id": user["id"], "nama": acrm["nama_acrm"], "role": "acrm"})
    await audit(user, "ACRM_APPROVE", "assets", asset_id)
    ma_user = await db.users.find_one({"ref_id": a["id_marketing_asset"], "role": "marketing_asset"})
    if ma_user:
        await notify(ma_user["id"], "asset", "Disetujui ACRM", f"Asset {a['nomor_asset']} disetujui ACRM, menunggu approval RCG.")
    for adm in await db.users.find({"role": "admin_rcg", "status": "active"}).to_list(10):
        await notify(adm["id"], "approval", "Menunggu approval RCG", f"Asset {a['nomor_asset']} menunggu approval Anda.")
    return {"ok": True, "status": new_status}

@api.post("/acrm/assets/{asset_id}/return")
async def acrm_return(asset_id: str, body: ReviewIn, user=Depends(require("acrm"))):
    if not body.notes or not body.notes.strip():
        raise HTTPException(400, "Catatan koreksi wajib diisi")
    acrm = await acrm_of(user)
    a = await db.assets.find_one({"id": asset_id, "deleted_at": None})
    if not a or a["id_acr"] != acrm["id_acr"]:
        raise HTTPException(403, "Anda tidak memiliki akses ke asset ini")
    if a["status"] not in ["WAITING_ACRM_REVIEW", "UPDATE_PENDING_ACRM"]:
        raise HTTPException(409, "Status asset sudah berubah, muat ulang data.")
    await db.assets.update_one({"id": asset_id}, {"$set": {"status": "RETURN_TO_MARKETING",
        "correction_notes": body.notes, "updated_at": now_iso()}})
    await approval_log(asset_id, "RETURN", a["status"], "RETURN_TO_MARKETING", {"id": user["id"], "nama": acrm["nama_acrm"], "role": "acrm"}, body.notes)
    await audit(user, "ACRM_RETURN", "assets", asset_id, notes=body.notes)
    ma_user = await db.users.find_one({"ref_id": a["id_marketing_asset"], "role": "marketing_asset"})
    if ma_user:
        await notify(ma_user["id"], "asset", "Dikembalikan ACRM", f"Asset {a['nomor_asset']} dikembalikan ACRM. Lihat catatan koreksi.")
    return {"ok": True}

# =================== RCG ===================
@api.get("/rcg/pending")
async def rcg_pending(user=Depends(require("admin_rcg"))):
    q = {"status": {"$in": ["WAITING_RCG_APPROVAL", "UPDATE_PENDING_RCG"]}, "deleted_at": None}
    return [await asset_internal_view(clean(a)) async for a in db.assets.find(q).sort("updated_at", 1)]

@api.post("/rcg/assets/{asset_id}/approve")
async def rcg_approve(asset_id: str, user=Depends(require("admin_rcg"))):
    a = await db.assets.find_one({"id": asset_id, "deleted_at": None})
    if not a:
        raise HTTPException(404, "Asset tidak ditemukan")
    if a["status"] not in ["WAITING_RCG_APPROVAL", "UPDATE_PENDING_RCG"]:
        raise HTTPException(409, "Status asset sudah berubah, muat ulang data.")
    # placeholder data cannot be published publicly
    ma_flag = (await db.master_marketing_asset.find_one({"id": a["id_marketing_asset"]}) or {}).get("data_flag")
    public_ready = ma_flag != "PERLU_KONFIRMASI_DATA"
    await db.assets.update_one({"id": asset_id}, {"$set": {"status": "PUBLISHED",
        "public_ready": public_ready, "published_at": now_iso(),
        "current_version_no": a.get("current_version_no", 1) + (1 if a["status"] == "UPDATE_PENDING_RCG" else 0),
        "updated_at": now_iso()}})
    await approval_log(asset_id, "PUBLISH", a["status"], "PUBLISHED", {"id": user["id"], "nama": user.get("nama"), "role": "admin_rcg"})
    await audit(user, "RCG_PUBLISH", "assets", asset_id, notes=None if public_ready else "Tidak tampil publik: data PIC placeholder")
    ma_user = await db.users.find_one({"ref_id": a["id_marketing_asset"], "role": "marketing_asset"})
    if ma_user:
        await notify(ma_user["id"], "asset", "Dipublikasikan", f"Asset {a['nomor_asset']} telah dipublikasikan ke katalog publik.")
    return {"ok": True, "status": "PUBLISHED", "public_ready": public_ready}

@api.post("/rcg/assets/{asset_id}/sold")
async def rcg_mark_sold(asset_id: str, body: Optional[ReviewIn] = None, user=Depends(require("admin_rcg"))):
    """Controller marks a published asset as SOLD: stays visible with TERJUAL badge, WA contact hidden."""
    a = await db.assets.find_one({"id": asset_id, "deleted_at": None})
    if not a:
        raise HTTPException(404, "Asset tidak ditemukan")
    if a["status"] != "PUBLISHED":
        raise HTTPException(409, "Hanya asset berstatus PUBLISHED yang dapat ditandai terjual")
    notes = (body.notes if body else None) or None
    await db.assets.update_one({"id": asset_id}, {"$set": {"status": "SOLD", "sold_at": now_iso(), "sold_notes": notes, "updated_at": now_iso()}})
    await approval_log(asset_id, "MARK_SOLD", "PUBLISHED", "SOLD", {"id": user["id"], "nama": user.get("nama"), "role": "admin_rcg"}, notes)
    await audit(user, "RCG_MARK_SOLD", "assets", asset_id, notes=notes)
    for role in ["marketing_asset", "acrm"]:
        ref = a["id_marketing_asset"] if role == "marketing_asset" else a["id_acrm"]
        u = await db.users.find_one({"ref_id": ref, "role": role})
        if u:
            await notify(u["id"], "asset", "Asset Terjual", f"Asset {a['nomor_asset']} ditandai TERJUAL oleh RCG. Kontak WhatsApp disembunyikan dari publik.")
    return {"ok": True, "status": "SOLD"}

@api.post("/rcg/assets/{asset_id}/unsold")
async def rcg_unmark_sold(asset_id: str, user=Depends(require("admin_rcg"))):
    """Revert an accidental SOLD mark back to PUBLISHED."""
    a = await db.assets.find_one({"id": asset_id, "deleted_at": None})
    if not a:
        raise HTTPException(404, "Asset tidak ditemukan")
    if a["status"] != "SOLD":
        raise HTTPException(409, "Asset tidak berstatus SOLD")
    await db.assets.update_one({"id": asset_id}, {"$set": {"status": "PUBLISHED", "sold_at": None, "sold_notes": None, "updated_at": now_iso()}})
    await approval_log(asset_id, "UNMARK_SOLD", "SOLD", "PUBLISHED", {"id": user["id"], "nama": user.get("nama"), "role": "admin_rcg"})
    await audit(user, "RCG_UNMARK_SOLD", "assets", asset_id)
    return {"ok": True, "status": "PUBLISHED"}

# ---------- Excel report (Admin RCG) ----------
REPORT_STATUSES = ["DRAFT", "WAITING_ACRM_REVIEW", "RETURN_TO_MARKETING", "WAITING_RCG_APPROVAL", "RETURN_FROM_RCG",
                   "PUBLISHED", "UPDATE_PENDING_ACRM", "UPDATE_PENDING_RCG", "SOLD"]

@api.get("/rcg/reports/assets/link")
async def rcg_report_link(user=Depends(require("admin_rcg"))):
    token = jwt.encode({"typ": "report", "uid": user["id"],
        "exp": datetime.now(timezone.utc) + timedelta(minutes=10)}, JWT_SECRET, algorithm=JWT_ALG)
    await audit(user, "EXPORT_REPORT", "assets", None)
    return {"url": f"/api/rcg/reports/assets.xlsx?token={token}", "expires_in": 600}

def build_report_xlsx(acrs, assets, cats, kpknls):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF"); head_fill = PatternFill("solid", fgColor="00A0A0")
    def header(ws, cols):
        ws.append(cols)
        for i, _ in enumerate(cols, 1):
            c = ws.cell(row=1, column=i); c.font = head_font; c.fill = head_fill; c.alignment = Alignment(horizontal="center")
    def autosize(ws):
        for i, col in enumerate(ws.columns, 1):
            width = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
    # Sheet 1: summary per ACR x status
    ws = wb.active; ws.title = "Ringkasan per ACR"
    header(ws, ["ACR", "Identifier"] + REPORT_STATUSES + ["Total", "Nilai Harga Limit (Published+Sold)"])
    by_acr = {}
    for a in assets:
        by_acr.setdefault(a["id_acr"], []).append(a)
    for acr in acrs:
        rows = by_acr.get(acr["id"], [])
        counts = [sum(1 for a in rows if a["status"] == s) for s in REPORT_STATUSES]
        value = sum((a.get("harga_limit") or 0) for a in rows if a["status"] in ("PUBLISHED", "SOLD"))
        ws.append([acr["nama_acr"], acr.get("identifier")] + counts + [len(rows), value])
    tot = ["TOTAL", ""] + [sum(1 for a in assets if a["status"] == s) for s in REPORT_STATUSES] + [len(assets), sum((a.get("harga_limit") or 0) for a in assets if a["status"] in ("PUBLISHED", "SOLD"))]
    ws.append(tot)
    for c in ws[ws.max_row]: c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=ws.max_column, max_col=ws.max_column):
        for c in row: c.number_format = '"Rp"#,##0'
    autosize(ws)
    # Sheet 2: detail
    ws2 = wb.create_sheet("Detail Asset")
    header(ws2, ["Nomor Asset", "Judul", "Status", "ACR", "ACRM", "PIC Marketing", "Kategori", "Subkategori", "Provinsi", "Kabupaten/Kota",
                 "Kecamatan", "Kelurahan/Desa", "Harga Limit", "Nilai Appraisal", "Luas Tanah", "Luas Bangunan", "Tanggal Lelang", "KPKNL",
                 "Dipublikasikan", "Terjual", "Dibuat"])
    for a in sorted(assets, key=lambda x: (x.get("acr_nama") or "", x.get("nomor_asset") or "")):
        sched = a.get("schedule") or {}
        ws2.append([a.get("nomor_asset"), a.get("judul_asset"), a.get("status"), a.get("acr_nama"), a.get("acrm_nama"), a.get("pic_nama"),
                    cats.get(a.get("id_category")), cats.get(a.get("id_subcategory")), a.get("provinsi"), a.get("kabupaten_kota"),
                    a.get("kecamatan"), a.get("wilayah_level_4"), a.get("harga_limit"), a.get("nilai_appraisal"), a.get("luas_tanah"), a.get("luas_bangunan"),
                    sched.get("tanggal_lelang"), kpknls.get(sched.get("id_kpknl")), (a.get("published_at") or "")[:10], (a.get("sold_at") or "")[:10], (a.get("created_at") or "")[:10]])
    for row in ws2.iter_rows(min_row=2, min_col=13, max_col=14):
        for c in row: c.number_format = '"Rp"#,##0'
    ws2.freeze_panes = "A2"; autosize(ws2)
    # Sheet 3: per status
    ws3 = wb.create_sheet("Ringkasan Status")
    header(ws3, ["Status", "Jumlah Asset", "Total Harga Limit"])
    for s in REPORT_STATUSES:
        rows = [a for a in assets if a["status"] == s]
        ws3.append([s, len(rows), sum((a.get("harga_limit") or 0) for a in rows)])
    for row in ws3.iter_rows(min_row=2, min_col=3, max_col=3):
        for c in row: c.number_format = '"Rp"#,##0'
    autosize(ws3)
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()

@api.get("/rcg/reports/assets.xlsx")
async def rcg_report_xlsx(token: str = Query(...)):
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALG])
    except jwt.PyJWTError:
        raise HTTPException(401, "Tautan tidak valid atau sudah kadaluarsa")
    if payload.get("typ") != "report":
        raise HTTPException(401, "Tautan tidak valid")
    u = await db.users.find_one({"id": payload.get("uid"), "role": "admin_rcg"})
    if not u:
        raise HTTPException(403, "Akses ditolak")
    acrs = [a async for a in db.master_acr.find({"deleted_at": None}).sort("nama_acr", 1)]
    assets = [a async for a in db.assets.find({"deleted_at": None})]
    cats = {c["id"]: c["nama_category"] async for c in db.master_asset_category.find({})}
    kpknls = {k["id"]: k["nama_kpknl"] async for k in db.master_kpknl.find({})}
    data = await run_in_threadpool(build_report_xlsx, acrs, assets, cats, kpknls)
    fname = f"Laporan_Asset_BSI_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return FastResponse(content=data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'})

@api.post("/rcg/assets/{asset_id}/return")
async def rcg_return(asset_id: str, body: ReviewIn, user=Depends(require("admin_rcg"))):
    if not body.notes or not body.notes.strip():
        raise HTTPException(400, "Catatan koreksi wajib diisi")
    a = await db.assets.find_one({"id": asset_id, "deleted_at": None})
    if not a:
        raise HTTPException(404, "Asset tidak ditemukan")
    if a["status"] not in ["WAITING_RCG_APPROVAL", "UPDATE_PENDING_RCG"]:
        raise HTTPException(409, "Status asset sudah berubah, muat ulang data.")
    await db.assets.update_one({"id": asset_id}, {"$set": {"status": "RETURN_FROM_RCG",
        "correction_notes": body.notes, "updated_at": now_iso()}})
    await approval_log(asset_id, "RETURN", a["status"], "RETURN_FROM_RCG", {"id": user["id"], "nama": user.get("nama"), "role": "admin_rcg"}, body.notes)
    await audit(user, "RCG_RETURN", "assets", asset_id, notes=body.notes)
    ma_user = await db.users.find_one({"ref_id": a["id_marketing_asset"], "role": "marketing_asset"})
    if ma_user:
        await notify(ma_user["id"], "asset", "Dikembalikan RCG", f"Asset {a['nomor_asset']} dikembalikan RCG. Lihat catatan koreksi.")
    acrm_user = await db.users.find_one({"ref_id": a.get("id_acrm"), "role": "acrm"})
    if acrm_user:
        await notify(acrm_user["id"], "asset", "Dikembalikan RCG", f"Asset {a['nomor_asset']} dikembalikan RCG.")
    return {"ok": True}

# =================== DASHBOARDS ===================
async def count_status(match, statuses):
    match = {**match, "deleted_at": None}
    out = {}
    for s in statuses:
        out[s] = await db.assets.count_documents({**match, "status": s})
    return out

@api.get("/dashboard/marketing")
async def dash_marketing(user=Depends(require("marketing_asset"))):
    ma, acr, acrm = await ma_context(user)
    m = {"id_marketing_asset": ma["id"]}
    total = await db.assets.count_documents({**m, "deleted_at": None})
    st = await count_status(m, ["DRAFT", "WAITING_ACRM_REVIEW", "RETURN_TO_MARKETING",
        "WAITING_RCG_APPROVAL", "RETURN_FROM_RCG", "PUBLISHED", "SOLD", "UPDATE_PENDING_ACRM", "UPDATE_PENDING_RCG"])
    return {"total": total, "acr": acr["nama_acr"] if acr else None, "by_status": st}

@api.get("/dashboard/acrm")
async def dash_acrm(user=Depends(require("acrm"))):
    acrm = await acrm_of(user)
    acr = await db.master_acr.find_one({"id": acrm["id_acr"]})
    m = {"id_acr": acrm["id_acr"]}
    total = await db.assets.count_documents({**m, "deleted_at": None})
    st = await count_status(m, ["WAITING_ACRM_REVIEW", "UPDATE_PENDING_ACRM", "WAITING_RCG_APPROVAL",
        "RETURN_TO_MARKETING", "PUBLISHED"])
    return {"total": total, "acr": acr["nama_acr"] if acr else None, "by_status": st}

@api.get("/dashboard/rcg")
async def dash_rcg(user=Depends(require("admin_rcg"))):
    total_acr = await db.master_acr.count_documents({"deleted_at": None})
    total_acrm = await db.master_acrm.count_documents({"deleted_at": None})
    total_ma = await db.master_marketing_asset.count_documents({"deleted_at": None})
    total_asset = await db.assets.count_documents({"deleted_at": None})
    st = await count_status({}, ["WAITING_ACRM_REVIEW", "WAITING_RCG_APPROVAL", "UPDATE_PENDING_RCG",
        "RETURN_TO_MARKETING", "RETURN_FROM_RCG", "PUBLISHED", "SOLD"])
    with_sched = await db.assets.count_documents({"deleted_at": None, "schedule.tanggal_lelang": {"$exists": True, "$ne": None}})
    # per ACR
    per_acr = []
    async for acr in db.master_acr.find({"deleted_at": None}).sort("nama_acr", 1):
        m = {"id_acr": acr["id"], "deleted_at": None}
        per_acr.append({"nama_acr": acr["nama_acr"],
            "total": await db.assets.count_documents(m),
            "published": await db.assets.count_documents({**m, "status": "PUBLISHED"}),
            "pending": await db.assets.count_documents({**m, "status": {"$in": ["WAITING_ACRM_REVIEW", "WAITING_RCG_APPROVAL", "UPDATE_PENDING_ACRM", "UPDATE_PENDING_RCG"]}}),
            "returned": await db.assets.count_documents({**m, "status": {"$in": ["RETURN_TO_MARKETING", "RETURN_FROM_RCG"]}})})
    return {"totals": {"acr": total_acr, "acrm": total_acrm, "marketing_asset": total_ma,
        "asset": total_asset, "with_schedule": with_sched}, "by_status": st, "per_acr": per_acr}

# =================== NOTIFICATIONS ===================
@api.get("/notifications")
async def get_notifications(user=Depends(current_user)):
    items = [clean(n) async for n in db.notifications.find({"user_id": user["id"]}).sort("created_at", -1).limit(50)]
    unread = await db.notifications.count_documents({"user_id": user["id"], "is_read": False})
    return {"items": items, "unread": unread}

@api.post("/notifications/{nid}/read")
async def read_notification(nid: str, user=Depends(current_user)):
    await db.notifications.update_one({"id": nid, "user_id": user["id"]}, {"$set": {"is_read": True}})
    return {"ok": True}

@api.post("/notifications/read-all")
async def read_all(user=Depends(current_user)):
    await db.notifications.update_many({"user_id": user["id"]}, {"$set": {"is_read": True}})
    return {"ok": True}

# =================== ADMIN: master + user mgmt ===================
@api.get("/admin/acr")
async def admin_acr(user=Depends(require("admin_rcg"))):
    return [clean(a) async for a in db.master_acr.find({"deleted_at": None}).sort("nama_acr", 1)]

class CategoryIn(BaseModel):
    nama_category: str
    parent_category_id: Optional[str] = None

@api.post("/admin/category")
async def add_category(body: CategoryIn, user=Depends(require("admin_rcg"))):
    c = {"id": new_id(), "nama_category": body.nama_category, "parent_category_id": body.parent_category_id,
         "status": "active", "created_at": now_iso(), "updated_at": now_iso(), "deleted_at": None}
    await db.master_asset_category.insert_one(c)
    await audit(user, "CREATE", "category", c["id"], after={"nama": body.nama_category})
    return clean(c)

@api.put("/admin/category/{cid}")
async def edit_category(cid: str, body: CategoryIn, user=Depends(require("admin_rcg"))):
    await db.master_asset_category.update_one({"id": cid}, {"$set": {"nama_category": body.nama_category, "updated_at": now_iso()}})
    await audit(user, "UPDATE", "category", cid)
    return clean(await db.master_asset_category.find_one({"id": cid}))

@api.post("/admin/category/{cid}/toggle")
async def toggle_category(cid: str, user=Depends(require("admin_rcg"))):
    c = await db.master_asset_category.find_one({"id": cid})
    if not c:
        raise HTTPException(404, "Kategori tidak ditemukan")
    ns = "inactive" if c["status"] == "active" else "active"
    await db.master_asset_category.update_one({"id": cid}, {"$set": {"status": ns, "updated_at": now_iso()}})
    await audit(user, "TOGGLE_STATUS", "category", cid, after={"status": ns})
    return {"ok": True, "status": ns}

class KpknlIn(BaseModel):
    nama_kpknl: str
    alamat_kpknl: str
    provinsi: Optional[str] = None

@api.get("/admin/kpknl")
async def admin_kpknl(user=Depends(require("admin_rcg"))):
    return [clean(k) async for k in db.master_kpknl.find({"deleted_at": None}).sort("nama_kpknl", 1)]

@api.post("/admin/kpknl")
async def add_kpknl(body: KpknlIn, user=Depends(require("admin_rcg"))):
    k = {"id": new_id(), **body.model_dump(), "status": "active", "created_at": now_iso(),
         "updated_at": now_iso(), "deleted_at": None}
    await db.master_kpknl.insert_one(k)
    await audit(user, "CREATE", "kpknl", k["id"], after={"nama": body.nama_kpknl})
    return clean(k)

@api.get("/admin/users")
async def admin_users(role: Optional[str] = None, keyword: Optional[str] = None, user=Depends(require("admin_rcg"))):
    coll = db.master_marketing_asset if role == "marketing_asset" else db.master_acrm if role == "acrm" else None
    result = []
    async def build(coll, r):
        async for u in coll.find({"deleted_at": None}).sort("created_at", 1):
            acr = await db.master_acr.find_one({"id": u["id_acr"]})
            nama = u.get("nama_marketing_asset") or u.get("nama_acrm")
            nip = u.get("nip") or u.get("nip_acrm")
            if keyword and keyword.lower() not in (nama or "").lower() and keyword.lower() not in (nip or "").lower():
                continue
            uu = await db.users.find_one({"ref_id": u["id"]})
            result.append({"id": u["id"], "nama": nama, "nip": nip, "role": r,
                "nomor_hp": u.get("nomor_hp"), "acr": acr["nama_acr"] if acr else None,
                "id_acr": u["id_acr"], "status": u["status"], "data_flag": u.get("data_flag"),
                "login_status": uu["status"] if uu else None, "user_id": uu["id"] if uu else None})
    if role in (None, "acrm"):
        await build(db.master_acrm, "acrm")
    if role in (None, "marketing_asset"):
        await build(db.master_marketing_asset, "marketing_asset")
    return result

class UserEditIn(BaseModel):
    nama: str
    nomor_hp: Optional[str] = None

def validate_hp(hp):
    d = re.sub(r'\D', '', hp or '')
    if d and (not d.startswith('0') or len(d) not in (10, 11, 12)):
        raise HTTPException(400, "Nomor HP harus digit, diawali 0, dan panjang 10, 11, atau 12 digit")
    return d

@api.put("/admin/marketing-asset/{mid}")
async def edit_ma(mid: str, body: UserEditIn, user=Depends(require("admin_rcg"))):
    m = await db.master_marketing_asset.find_one({"id": mid, "deleted_at": None})
    if not m:
        raise HTTPException(404, "Data tidak ditemukan")
    validate_hp(body.nomor_hp)
    before = {"nama": m["nama_marketing_asset"], "hp": m["nomor_hp"]}
    flag = m.get("data_flag")
    if flag == "PERLU_KONFIRMASI_DATA" and body.nama and body.nama != "NOPE":
        flag = "OK"
    await db.master_marketing_asset.update_one({"id": mid}, {"$set": {
        "nama_marketing_asset": body.nama, "nomor_hp": body.nomor_hp or m["nomor_hp"],
        "data_flag": flag, "updated_at": now_iso()}})
    await db.users.update_one({"ref_id": mid}, {"$set": {"nama": body.nama, "data_flag": flag}})
    await audit(user, "UPDATE", "marketing_asset", mid, before=before, after={"nama": body.nama, "hp": body.nomor_hp})
    return {"ok": True}

@api.put("/admin/acrm/{aid}")
async def edit_acrm(aid: str, body: UserEditIn, user=Depends(require("admin_rcg"))):
    m = await db.master_acrm.find_one({"id": aid, "deleted_at": None})
    if not m:
        raise HTTPException(404, "Data tidak ditemukan")
    before = {"nama": m["nama_acrm"]}
    flag = m.get("data_flag")
    if flag == "PERLU_KONFIRMASI_DATA" and body.nama and body.nama != "NOPE":
        flag = "OK"
    await db.master_acrm.update_one({"id": aid}, {"$set": {"nama_acrm": body.nama, "data_flag": flag, "updated_at": now_iso()}})
    await db.users.update_one({"ref_id": aid}, {"$set": {"nama": body.nama, "data_flag": flag}})
    await audit(user, "UPDATE", "acrm", aid, before=before, after={"nama": body.nama})
    return {"ok": True}

@api.post("/admin/{kind}/{uid}/toggle")
async def toggle_user(kind: str, uid: str, user=Depends(require("admin_rcg"))):
    coll = db.master_marketing_asset if kind == "marketing-asset" else db.master_acrm
    m = await coll.find_one({"id": uid, "deleted_at": None})
    if not m:
        raise HTTPException(404, "Data tidak ditemukan")
    ns = "inactive" if m["status"] == "active" else "active"
    await coll.update_one({"id": uid}, {"$set": {"status": ns, "updated_at": now_iso()}})
    await db.users.update_one({"ref_id": uid}, {"$set": {"status": ns}})
    await audit(user, "TOGGLE_STATUS", kind, uid, after={"status": ns})
    return {"ok": True, "status": ns}

class MutasiIn(BaseModel):
    new_acr_id: str
    reason: str

@api.post("/admin/marketing-asset/{mid}/mutasi")
async def mutasi_ma(mid: str, body: MutasiIn, user=Depends(require("admin_rcg"))):
    if not body.reason.strip():
        raise HTTPException(400, "Alasan mutasi wajib diisi")
    m = await db.master_marketing_asset.find_one({"id": mid, "deleted_at": None})
    if not m:
        raise HTTPException(404, "Data tidak ditemukan")
    new_acr = await db.master_acr.find_one({"id": body.new_acr_id, "status": "active", "deleted_at": None})
    if not new_acr:
        raise HTTPException(400, "ACR tujuan tidak valid")
    old_acr = m["id_acr"]
    await db.master_marketing_asset.update_one({"id": mid}, {"$set": {"id_acr": body.new_acr_id, "updated_at": now_iso()}})
    await db.user_mutation_history.insert_one({"id": new_id(), "user_id": mid, "user_role": "marketing_asset",
        "old_acr_id": old_acr, "new_acr_id": body.new_acr_id, "mutation_date": now_iso(),
        "mutation_by": user["username"], "reason": body.reason})
    await audit(user, "MUTASI", "marketing_asset", mid, before={"acr": old_acr}, after={"acr": body.new_acr_id}, notes=body.reason)
    uu = await db.users.find_one({"ref_id": mid})
    if uu:
        await notify(uu["id"], "mutasi", "Mutasi ACR", f"Anda dipindahkan ke {new_acr['nama_acr']}.")
    return {"ok": True}

@api.post("/admin/acrm/{aid}/mutasi")
async def mutasi_acrm(aid: str, body: MutasiIn, user=Depends(require("admin_rcg"))):
    if not body.reason.strip():
        raise HTTPException(400, "Alasan mutasi wajib diisi")
    m = await db.master_acrm.find_one({"id": aid, "deleted_at": None})
    if not m:
        raise HTTPException(404, "Data tidak ditemukan")
    new_acr = await db.master_acr.find_one({"id": body.new_acr_id, "status": "active", "deleted_at": None})
    if not new_acr:
        raise HTTPException(400, "ACR tujuan tidak valid")
    existing = await db.master_acrm.find_one({"id_acr": body.new_acr_id, "status": "active", "deleted_at": None, "id": {"$ne": aid}})
    if existing:
        raise HTTPException(400, f"ACR tujuan sudah memiliki ACRM aktif ({existing['nama_acrm']})")
    old_acr = m["id_acr"]
    await db.master_acrm.update_one({"id": aid}, {"$set": {"id_acr": body.new_acr_id, "updated_at": now_iso()}})
    await db.user_mutation_history.insert_one({"id": new_id(), "user_id": aid, "user_role": "acrm",
        "old_acr_id": old_acr, "new_acr_id": body.new_acr_id, "mutation_date": now_iso(),
        "mutation_by": user["username"], "reason": body.reason})
    await audit(user, "MUTASI", "acrm", aid, before={"acr": old_acr}, after={"acr": body.new_acr_id}, notes=body.reason)
    return {"ok": True}

@api.get("/admin/audit-logs")
async def get_audit(limit: int = 100, user=Depends(require("admin_rcg"))):
    return [clean(l) async for l in db.audit_logs.find().sort("timestamp", -1).limit(limit)]

@api.get("/admin/mutation-history")
async def get_mutations(user=Depends(require("admin_rcg"))):
    out = []
    async for h in db.user_mutation_history.find().sort("mutation_date", -1).limit(100):
        old = await db.master_acr.find_one({"id": h["old_acr_id"]})
        new = await db.master_acr.find_one({"id": h["new_acr_id"]})
        h = clean(h)
        h["old_acr"] = old["nama_acr"] if old else None
        h["new_acr"] = new["nama_acr"] if new else None
        out.append(h)
    return out

app.include_router(api)
app.add_middleware(CORSMiddleware, allow_credentials=True, allow_origins=["*"],
                   allow_methods=["*"], allow_headers=["*"])

@app.on_event("shutdown")
async def shutdown():
    client.close()
